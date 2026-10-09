from flask import Blueprint, request, send_file
import requests
import os
import re
import fitz
from pptx import Presentation
from pptx.util import Pt, Inches
import nltk
from bs4 import BeautifulSoup

nltk.download('punkt')

research_bp = Blueprint("research_bp", __name__)

UPLOAD_FOLDER = "uploads"
IMAGE_FOLDER = "uploads/images"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(IMAGE_FOLDER, exist_ok=True)

# ----------------------------------
# VALID SECTION LIST
# ----------------------------------

VALID_SECTIONS = [
    "abstract", "introduction", "background",
    "related work", "literature review",
    "methodology", "methods", "materials and methods",
    "proposed", "implementation",
    "experiment", "experimental setup",
    "results", "discussion",
    "results and discussion",
    "analysis", "evaluation",
    "conclusion", "future work"
]

INVALID_SECTIONS = [
    "author", "affiliation", "email", "doi",
    "citation", "publisher", "license",
    "access", "copyright",
    "received", "accepted", "published",
    "acknowledgment", "acknowledgement",
    "reference", "bibliography"
]

# ----------------------------------
# CLEAN TEXT
# ----------------------------------

def clean_text(text):

    text = re.sub(r'\S+@\S+', '', text)
    text = re.sub(r'\[\d+(,\s*\d+)*\]', '', text)
    text = re.sub(r'\(.*?\d{4}.*?\)', '', text)
    text = re.sub(r'\(\s*\)', '', text)
    text = re.sub(r'\s+', ' ', text)

    return text.strip()

# ----------------------------------
# VALIDATE HEADING
# ----------------------------------

def is_valid_heading(text):

    lower = text.lower().strip()

    for bad in INVALID_SECTIONS:
        if bad in lower:
            return False

    for good in VALID_SECTIONS:
        if good in lower:
            return True

    if re.match(r'^\d+\.?\s*(introduction|method|result|discussion|conclusion)', lower):
        return True

    return False


def clean_heading(title):

    title = re.sub(r'^\d+\.?\s*', '', title)
    title = re.sub(r'\(.*?\)', '', title)
    title = re.sub(r'\[.*?\]', '', title)

    return title.strip()

def generate_short_title(title):

    title = title.lower()

    keyword_map = {
        "abstract": "Abstract",
        "introduction": "Introduction",
        "background": "Background",
        "related": "Related Work",
        "literature": "Literature",
        "method": "Method",
        "methodology": "Method",
        "proposed": "Method",
        "implementation": "Implementation",
        "experiment": "Experiment",
        "experimental": "Experiment",
        "result": "Results",
        "analysis": "Analysis",
        "discussion": "Discussion",
        "conclusion": "Conclusion",
        "future": "Future Work",
        "evaluation": "Evaluation"
    }

    for key in keyword_map:

        if key in title:
            return keyword_map[key]

    # fallback: return first 2 words
    words = title.split()

    return " ".join(words[:2]).title()

# ----------------------------------
# CHECK LINK TYPE
# ----------------------------------

def is_pdf_link(url):
    return url.lower().endswith(".pdf")

# ----------------------------------
# DOWNLOAD PDF
# ----------------------------------

def download_pdf(url):

    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(url, headers=headers, timeout=30)

    if response.status_code != 200:
        return None

    path = os.path.join(UPLOAD_FOLDER, "paper.pdf")

    with open(path, "wb") as f:
        f.write(response.content)

    return path

# ----------------------------------
# EXTRACT FROM PDF WITH IMAGE MATCHING
# ----------------------------------

def extract_from_pdf(pdf_path):

    doc = fitz.open(pdf_path)

    sections = []
    current_section = None
    image_index = 0

    font_sizes = []

    # detect body font size
    for page in doc:
        blocks = page.get_text("dict")["blocks"]

        for block in blocks:
            if "lines" not in block:
                continue

            for line in block["lines"]:
                for span in line["spans"]:
                    font_sizes.append(span["size"])

    body_size = sorted(font_sizes)[len(font_sizes)//2]

    stop = False

    for page in doc:

        if stop:
            break

        blocks = page.get_text("dict")["blocks"]

        # sort blocks top to bottom
        blocks = sorted(blocks, key=lambda b: b["bbox"][1])

        section_positions = []

        for block in blocks:

            if "lines" not in block:
                continue

            text = ""
            max_size = 0

            for line in block["lines"]:
                for span in line["spans"]:

                    t = clean_text(span["text"])

                    if not t:
                        continue

                    text += t + " "

                    if span["size"] > max_size:
                        max_size = span["size"]

            text = text.strip()

            if not text:
                continue

            if "reference" in text.lower():
                stop = True
                break

            y_pos = block["bbox"][1]

            is_heading = False

            if max_size > body_size + 0.7:
                is_heading = True

            if text.isupper() and len(text) < 80:
                is_heading = True

            if re.match(r'^\d+\.?\s+', text):
                is_heading = True

            if is_heading and is_valid_heading(text):

                current_section = {
                    "title": generate_short_title(clean_heading(text)),
                    "content": "",
                    "images": [],
                    "y": y_pos
                    }


                sections.append(current_section)
                section_positions.append(current_section)

            else:

                if current_section and len(text.split()) > 5:
                    current_section["content"] += text + " "

        # Extract images and match to section
        for img in page.get_images(full=True):

            xref = img[0]

            pix = fitz.Pixmap(doc, xref)

            if pix.n < 5:

                img_path = f"{IMAGE_FOLDER}/img_{image_index}.png"
                pix.save(img_path)

                bbox = page.get_image_bbox(img)
                img_y = bbox.y0

                closest = None
                min_dist = 999999

                for sec in section_positions:

                    dist = img_y - sec["y"]

                    if dist >= 0 and dist < min_dist:
                        min_dist = dist
                        closest = sec

                if closest:
                    closest["images"].append(img_path)

                image_index += 1

            pix = None

    return sections

# ----------------------------------
# EXTRACT FROM HTML
# ----------------------------------

# ----------------------------------
# EXTRACT FROM HTML WITH IMAGE MATCHING
# ----------------------------------

def extract_from_html(url):

    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(url, headers=headers)

    soup = BeautifulSoup(response.text, "lxml")

    sections = []
    current_section = None
    img_index = 0

    # iterate in order: headings, paragraphs, images
    for element in soup.find_all(["h1","h2","h3","h4","p","img","figure"]):

        # ------------------
        # HEADING
        # ------------------
        if element.name.startswith("h"):

            text = clean_text(element.get_text())

            if is_valid_heading(text):

                current_section = {
                    "title": generate_short_title(clean_heading(text)),
                    "content": "",
                    "images": []
                }

                sections.append(current_section)

        # ------------------
        # PARAGRAPH
        # ------------------
        elif element.name == "p":

            if current_section:

                text = clean_text(element.get_text())

                if len(text.split()) > 5:
                    current_section["content"] += text + " "

        # ------------------
        # IMAGE TAG
        # ------------------
        elif element.name == "img":

            if current_section:

                src = element.get("src")

                if not src:
                    continue

                # fix relative URLs
                if src.startswith("//"):
                    src = "https:" + src

                elif src.startswith("/"):
                    base = "/".join(url.split("/")[:3])
                    src = base + src

                try:

                    img_data = requests.get(src, headers=headers).content

                    img_path = f"{IMAGE_FOLDER}/html_img_{img_index}.png"

                    with open(img_path, "wb") as f:
                        f.write(img_data)

                    current_section["images"].append(img_path)

                    img_index += 1

                except:
                    pass

        # ------------------
        # FIGURE TAG (workflow diagrams usually here)
        # ------------------
        elif element.name == "figure":

            if current_section:

                img_tag = element.find("img")

                if img_tag:

                    src = img_tag.get("src")

                    if src:

                        if src.startswith("//"):
                            src = "https:" + src

                        elif src.startswith("/"):
                            base = "/".join(url.split("/")[:3])
                            src = base + src

                        try:

                            img_data = requests.get(src, headers=headers).content

                            img_path = f"{IMAGE_FOLDER}/html_img_{img_index}.png"

                            with open(img_path, "wb") as f:
                                f.write(img_data)

                            current_section["images"].append(img_path)

                            img_index += 1

                        except:
                            pass

    return sections

# ----------------------------------
# SUMMARIZE
# ----------------------------------

def summarize(text, max_points):

    sentences = nltk.sent_tokenize(text)

    bullets = []
    seen = set()

    for s in sentences:

        s = clean_text(s)

        if len(s.split()) < 8 or len(s.split()) > 30:
            continue

        if s in seen:
            continue

        seen.add(s)

        bullets.append(s)

        if len(bullets) >= max_points:
            break

    return bullets

# ----------------------------------
# CREATE PPT WITH IMAGE MATCHING
# ----------------------------------

def create_ppt(title, owner, sections, font_size, max_points, max_images=4):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()

    # -----------------------------
    # Cover Slide (dynamic sizing)
    # -----------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank slide
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    # Title textbox
    max_width = slide_width - Inches(1)  # 0.5 inch margin each side
    title_left = Inches(0.5)
    title_top = Inches(1.0)
    title_height = Inches(2)

    title_box = slide.shapes.add_textbox(title_left, title_top, max_width, title_height)
    tf_title = title_box.text_frame
    tf_title.clear()
    p_title = tf_title.add_paragraph()
    p_title.text = title
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.alignment = PP_ALIGN.CENTER
    tf_title.word_wrap = True

    # Presented By textbox
    owner_width = slide_width - Inches(2)
    owner_height = Inches(1)
    owner_left = Inches(1)
    owner_top = title_top + title_height + Inches(0.5)

    owner_box = slide.shapes.add_textbox(owner_left, owner_top, owner_width, owner_height)
    tf_owner = owner_box.text_frame
    tf_owner.clear()
    p_owner = tf_owner.add_paragraph()
    p_owner.text = f"Presented By: {owner}"
    p_owner.font.size = Pt(20)
    p_owner.font.bold = True
    p_owner.alignment = PP_ALIGN.CENTER

    # -----------------------------
    # Create text slides
    # -----------------------------
    figure_counter = 1
    images_to_add = []

    for sec in sections:
        bullets = summarize(sec.get("content", ""), max_points)

        if bullets:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            if slide.shapes.title:
                slide.shapes.title.text = sec["title"]

            if len(slide.placeholders) > 1 and slide.placeholders[1]:
                tf = slide.placeholders[1].text_frame
                tf.clear()
            else:
                text_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
                tf = text_box.text_frame
                tf.clear()

            for bullet in bullets:
                p = tf.add_paragraph()
                p.text = bullet
                p.font.size = Pt(int(font_size))
                p.level = 0

        # Collect first few images for the end
        if "images" in sec and sec["images"]:
            for img_path in sec["images"]:
                if len(images_to_add) < max_images:
                    images_to_add.append((img_path, sec["title"]))
                else:
                    break

    # -----------------------------
    # Add selected images at the end
    # -----------------------------
    for img_path, sec_title in images_to_add:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        if slide.shapes.title:
            slide.shapes.title.text = f"Figure {figure_counter}: {sec_title}"
        else:
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
            tf_title = title_box.text_frame
            tf_title.text = f"Figure {figure_counter}: {sec_title}"
            for p in tf_title.paragraphs:
                p.font.size = Pt(24)
                p.font.bold = True

        figure_counter += 1

        # Add the image
        left = Inches(1)
        top = Inches(1.5)
        width = Inches(8)
        try:
            slide.shapes.add_picture(img_path, left, top, width=width)
        except:
            pass

    # -----------------------------
    # Add Thank You slide (centered)
    # -----------------------------
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text_width = Inches(6)
    text_height = Inches(1.5)
    left = (slide_width - text_width) / 2
    top = (slide_height - text_height) / 2

    title_box = slide.shapes.add_textbox(left, top, text_width, text_height)
    tf = title_box.text_frame
    tf.clear()
    p = tf.add_paragraph()
    p.text = "THANK YOU"
    p.font.size = Pt(60)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # -----------------------------
    # Save PPT
    # -----------------------------
    ppt_path = os.path.join(UPLOAD_FOLDER, "Research_Presentation.pptx")
    prs.save(ppt_path)
    return ppt_path


##new changes function for the 1st slide title
# ----------------------------------
# EXTRACT TITLE FROM PDF
# ----------------------------------
def extract_pdf_title(pdf_path):
    import fitz
    doc = fitz.open(pdf_path)
    page = doc[0]  # first page
    font_sizes = []
    texts = []

    blocks = page.get_text("dict")["blocks"]
    for block in blocks:
        if "lines" not in block:
            continue
        block_text = ""
        max_size = 0
        for line in block["lines"]:
            for span in line["spans"]:
                t = clean_text(span["text"])
                if not t:
                    continue
                block_text += t + " "
                if span["size"] > max_size:
                    max_size = span["size"]
        block_text = block_text.strip()
        if block_text:
            texts.append((block_text, max_size))
            font_sizes.append(max_size)

    # Take the text with the largest font as title
    if not texts:
        return "Research Presentation"

    max_font = max(font_sizes)
    for t, size in texts:
        if size == max_font and len(t.split()) < 25:  # avoid huge blocks
            return t
    return texts[0][0]

# ----------------------------------
# EXTRACT TITLE FROM HTML
# ----------------------------------
def extract_html_title(url):
    import requests
    from bs4 import BeautifulSoup
    headers = {"User-Agent": "Mozilla/5.0"}
    response = requests.get(url, headers=headers)
    soup = BeautifulSoup(response.text, "lxml")

    # try <title>
    if soup.title and soup.title.string:
        return clean_text(soup.title.string)

    # fallback to first <h1>
    h1 = soup.find("h1")
    if h1:
        return clean_text(h1.get_text())

    return "Research Presentation"

# ----------------------------------
# MAIN ROUTE
# ----------------------------------

@research_bp.route("/generate/research", methods=["POST"])
def generate_research_ppt():

    owner = request.form["owner_name"]
    url = request.form["url"]
    font_size = request.form["font_size"]
    max_points = int(request.form["max_points"])

    # -----------------------------
    # Extract title automatically
    # -----------------------------
    title = "Research Presentation"  # fallback title

    if is_pdf_link(url):
        pdf_path = download_pdf(url)
        if not pdf_path:
            return "PDF download failed"

        # Extract the title from PDF
        title_extracted = extract_pdf_title(pdf_path)
        if title_extracted:
            title = title_extracted

        sections = extract_from_pdf(pdf_path)

    else:
        # Extract the title from HTML page
        title_extracted = extract_html_title(url)
        if title_extracted:
            title = title_extracted

        sections = extract_from_html(url)

    if not sections:
        return "No valid research sections found"

    ppt = create_ppt(
        title,
        owner,
        sections,
        font_size,
        max_points
    )

    return send_file(ppt, as_attachment=True)
